#!/usr/bin/env python3
"""
Create PowerPoint analyzing first 50 subtitles with waveforms
Shows original timing, waveform, and final timing
"""

import json
import numpy as np
import librosa
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import re
import sys
from io import BytesIO

def parse_srt(srt_path, max_subtitles=50):
    """Parse SRT file and return first N subtitles"""
    with open(srt_path, 'r', encoding='utf-8') as f:
        content = f.read()

    blocks = content.strip().split('\n\n')
    subtitles = []

    for block in blocks[:max_subtitles]:
        lines = block.split('\n')
        if len(lines) >= 3:
            num = int(lines[0])
            timing = lines[1]
            text = '\n'.join(lines[2:])

            # Parse timing
            times = timing.split(' --> ')
            start_str, end_str = times[0], times[1]

            def to_seconds(t):
                parts = t.replace(',', ':').split(':')
                return int(parts[0])*3600 + int(parts[1])*60 + int(parts[2]) + float(parts[3])/1000

            start_sec = to_seconds(start_str)
            end_sec = to_seconds(end_str)
            duration = end_sec - start_sec

            subtitles.append({
                'num': num,
                'start': start_sec,
                'end': end_sec,
                'duration': duration,
                'text': text,
                'start_str': start_str,
                'end_str': end_str
            })

    return subtitles


def load_json_words(json_path):
    """Load original words from JSON"""
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    words = []
    for seg in data['segments']:
        for w in seg.get('words', []):
            text = w.get('word', '').strip()
            if text:
                words.append({
                    'text': text,
                    'start': float(w.get('start', 0)),
                    'end': float(w.get('end', 0))
                })

    return words


def find_original_words(subtitle_text, all_words, subtitle_start, subtitle_end):
    """Find original words that match this subtitle"""
    # Extract words from subtitle text (remove hyphens)
    sub_words = subtitle_text.replace('-', '').replace('\n', ' ').split()
    sub_words = [w.strip() for w in sub_words if w.strip()]

    # Find matching words in original JSON near this time
    matches = []
    for word in all_words:
        # Check if word is near subtitle time (within 5 seconds)
        if abs(word['start'] - subtitle_start) < 5 or abs(word['start'] - subtitle_end) < 5:
            if word['text'] in sub_words:
                matches.append(word)

    return matches


def create_waveform_image(audio, sr, start_time, end_time, original_words, subtitle_start, subtitle_end):
    """Create waveform visualization"""
    # Add buffer for context
    buffer = 0.5
    vis_start = max(0, start_time - buffer)
    vis_end = end_time + buffer

    start_sample = int(vis_start * sr)
    end_sample = int(vis_end * sr)
    segment = audio[start_sample:end_sample]

    # Time array
    times = vis_start + np.arange(len(segment)) / sr

    # Create figure
    fig, ax = plt.subplots(figsize=(10, 3))

    # Plot waveform
    ax.plot(times, segment, linewidth=0.5, color='gray', alpha=0.7)
    ax.fill_between(times, segment, color='lightgray', alpha=0.3)

    # Mark original word boundaries (BLUE)
    for word in original_words:
        ax.axvspan(word['start'], word['end'], alpha=0.15, color='blue', label='Original')
        ax.axvline(word['start'], color='blue', linestyle='--', linewidth=1, alpha=0.5)
        ax.axvline(word['end'], color='blue', linestyle='--', linewidth=1, alpha=0.5)

    # Mark final subtitle boundaries (GREEN)
    ax.axvspan(subtitle_start, subtitle_end, alpha=0.2, color='green', label='Final Subtitle')
    ax.axvline(subtitle_start, color='green', linestyle='-', linewidth=2, alpha=0.8)
    ax.axvline(subtitle_end, color='green', linestyle='-', linewidth=2, alpha=0.8)

    # Formatting
    ax.set_xlim(vis_start, vis_end)
    ax.set_xlabel('Time (seconds)', fontsize=10)
    ax.set_ylabel('Amplitude', fontsize=10)
    ax.set_title(f'Waveform: {vis_start:.2f}s - {vis_end:.2f}s', fontsize=11)
    ax.grid(True, alpha=0.3)

    # Legend (avoid duplicates)
    handles, labels = ax.get_legend_handles_labels()
    by_label = dict(zip(labels, handles))
    ax.legend(by_label.values(), by_label.keys(), loc='upper right', fontsize=8)

    # Save to bytes
    buf = BytesIO()
    plt.tight_layout()
    plt.savefig(buf, format='png', dpi=100, bbox_inches='tight')
    plt.close()
    buf.seek(0)

    return buf


def create_presentation(srt_path, json_path, audio_path, output_path, max_subtitles=50):
    """Create PowerPoint presentation"""

    print("Loading subtitle data...")
    subtitles = parse_srt(srt_path, max_subtitles)

    print("Loading original JSON words...")
    all_words = load_json_words(json_path)

    print("Loading audio...")
    audio, sr = librosa.load(audio_path, sr=16000, mono=True)

    print("Creating presentation...")
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "Subtitle Timing Analysis"
    subtitle.text = f"First {len(subtitles)} Subtitles\nOriginal vs Energy-Corrected Timing"

    # Content slides
    blank_layout = prs.slide_layouts[6]

    for i, sub in enumerate(subtitles):
        print(f"Processing subtitle {i+1}/{len(subtitles)}...")

        slide = prs.slides.add_slide(blank_layout)

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.5))
        title_frame = title_box.text_frame
        title_frame.text = f"Subtitle #{sub['num']}: {sub['start_str']} → {sub['end_str']}"
        title_frame.paragraphs[0].font.size = Pt(18)
        title_frame.paragraphs[0].font.bold = True

        # Subtitle text
        text_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(9), Inches(1.0))
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        text_frame.text = sub['text']
        text_frame.paragraphs[0].font.size = Pt(14)
        text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Find original words
        original_words = find_original_words(sub['text'], all_words, sub['start'], sub['end'])

        # Calculate original timing span
        if original_words:
            orig_start = min(w['start'] for w in original_words)
            orig_end = max(w['end'] for w in original_words)
            orig_duration = orig_end - orig_start
        else:
            orig_start = sub['start']
            orig_end = sub['end']
            orig_duration = sub['duration']

        # Timing info
        info_y = 2.0

        # Original timing (BLUE)
        orig_box = slide.shapes.add_textbox(Inches(0.5), Inches(info_y), Inches(4), Inches(0.6))
        orig_frame = orig_box.text_frame
        orig_frame.text = f"Original (WhisperX):\n{orig_start:.3f}s → {orig_end:.3f}s ({orig_duration:.3f}s)"
        orig_frame.paragraphs[0].font.size = Pt(11)
        orig_frame.paragraphs[0].font.color.rgb = RGBColor(0, 0, 255)

        # Final timing (GREEN)
        final_box = slide.shapes.add_textbox(Inches(5), Inches(info_y), Inches(4), Inches(0.6))
        final_frame = final_box.text_frame
        final_frame.text = f"Final (Energy-Corrected):\n{sub['start']:.3f}s → {sub['end']:.3f}s ({sub['duration']:.3f}s)"
        final_frame.paragraphs[0].font.size = Pt(11)
        final_frame.paragraphs[0].font.color.rgb = RGBColor(0, 128, 0)

        # Timing shift
        start_shift = sub['start'] - orig_start
        end_shift = sub['end'] - orig_end
        duration_change = sub['duration'] - orig_duration

        shift_box = slide.shapes.add_textbox(Inches(0.5), Inches(info_y + 0.7), Inches(9), Inches(0.4))
        shift_frame = shift_box.text_frame
        shift_frame.text = f"Shift: Start {start_shift:+.3f}s | End {end_shift:+.3f}s | Duration {duration_change:+.3f}s"
        shift_frame.paragraphs[0].font.size = Pt(10)
        shift_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Flag issues
        issues = []
        word_count = len(sub['text'].replace('-', '').split())
        if word_count <= 2 and sub['duration'] < 1.0:
            issues.append(f"⚠️ Very short subtitle ({word_count} word(s), {sub['duration']:.2f}s)")
        if sub['duration'] > 10.0:
            issues.append(f"⚠️ Very long subtitle ({sub['duration']:.2f}s)")
        if abs(start_shift) > 1.0:
            issues.append(f"⚠️ Large start shift ({start_shift:+.2f}s)")

        if issues:
            issue_box = slide.shapes.add_textbox(Inches(0.5), Inches(info_y + 1.2), Inches(9), Inches(0.4))
            issue_frame = issue_box.text_frame
            issue_frame.text = " | ".join(issues)
            issue_frame.paragraphs[0].font.size = Pt(10)
            issue_frame.paragraphs[0].font.color.rgb = RGBColor(255, 0, 0)
            issue_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Waveform
        waveform_start = min(orig_start, sub['start'])
        waveform_end = max(orig_end, sub['end'])

        img_buf = create_waveform_image(
            audio, sr,
            waveform_start, waveform_end,
            original_words,
            sub['start'], sub['end']
        )

        pic = slide.shapes.add_picture(img_buf, Inches(0.5), Inches(4), width=Inches(9))

    print(f"\nSaving presentation to {output_path}...")
    prs.save(output_path)
    print("✅ Done!")


if __name__ == '__main__':
    srt_path = r'd:\stevner\2025\regionalt\Lørdag\04-LørdagEttermiddag\04Transcribe.energy\04-LørdagEttermiddag_clean.srt'
    json_path = r'd:\stevner\2025\regionalt\Lørdag\04-LørdagEttermiddag\04-LørdagEttermiddag.json'
    audio_path = r'd:\stevner\2025\regionalt\Lørdag\04-LørdagEttermiddag\04-LørdagEttermiddag.mp3'
    output_path = r'd:\stevner\2025\regionalt\Lørdag\04-LørdagEttermiddag\first_50_subtitles_analysis.pptx'

    create_presentation(srt_path, json_path, audio_path, output_path, max_subtitles=50)
